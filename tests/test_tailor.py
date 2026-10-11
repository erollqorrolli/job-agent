from docx import Document
from jobagent.tailor import tailor, matched_skills

BANK = {"Python": "used", "Docker": "used", "Kubernetes": "course", "Go": "course", "Rust": "course", "Terraform": "course", "Redis": "used",
        "Distributed Systems": "course", "Java": "used", "Machine Learning": "used", "FastAPI": "used", "Solidity": "course"}


def make(path):
    d = Document()
    def para(parts):
        p = d.add_paragraph()
        for text, bold in parts:
            r = p.add_run(text); r.bold = bold or None
        return p
    para([("PROJECTS", None)])
    para([("Graph tool", True), (" | Python", None)]); para([("• Built a thing with Python and Docker", None)])
    para([("Agent tool", True), (" | Python, FastAPI", None)]); para([("• Built agents with Machine Learning and FastAPI and Rust", None)])
    para([("EXPERIENCE", None)]); para([("Job", True)])
    para([("TECHNICAL SKILLS", None)])
    para([("Programming Languages", True), (": Python, Java", None)])
    para([("Frameworks & Tools", True), (": FastAPI, Docker", None), (" ", None), ("Interests", True), (": AI/ML", None)])
    d.save(str(path))


def test_adds_only_confirmed_mentioned_skills(tmp_path):
    src = tmp_path / "r.docx"; make(src)
    info = tailor(src, "We use Go, Kubernetes, Terraform, Redis and Docker. Experience with distributed systems. Cobol a plus.", BANK, tmp_path / "o.docx")
    txt = "\n".join(p.text for p in Document(str(tmp_path / "o.docx")).paragraphs)
    assert "Redis" in txt                                    # used + mentioned -> tools line
    assert "Coursework & Familiarity: " in txt and "Kubernetes" in txt and "Distributed Systems" in txt
    skills = "\n".join(l for l in txt.splitlines() if l.startswith(("Programming", "Frameworks", "Coursework")))
    assert "Solidity" not in skills and "Rust" not in skills and "Cobol" not in skills   # not mentioned / not in bank
    tools = [p.text for p in Document(str(tmp_path / "o.docx")).paragraphs if p.text.startswith("Frameworks")][0]
    assert tools.index("Redis") < tools.index("Interests")  # inserted before the Interests label
    assert "Kubernetes" not in tools                        # course-level never on the tools line
    assert "Go" in info["added_course"]                     # golang/"Go," context accepted as listed


def test_ambiguous_words_do_not_match():
    assert ("Go", "course") not in matched_skills("Ready to go the extra mile", {"Go": "course"})


def test_projects_reorder_by_relevance_and_bullets_untouched(tmp_path):
    src = tmp_path / "r.docx"; make(src)
    info = tailor(src, "Machine Learning, FastAPI and Rust services", BANK, tmp_path / "o.docx")
    texts = [p.text for p in Document(str(tmp_path / "o.docx")).paragraphs]
    assert info["reordered"] and texts.index("Agent tool | Python, FastAPI") < texts.index("Graph tool | Python")
    assert "• Built a thing with Python and Docker" in texts
