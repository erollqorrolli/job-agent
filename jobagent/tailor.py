"""Per-job resume tailoring. Deterministic and conservative:
 - never rewrites or invents bullets
 - adds only skills from the user's confirmed bank that the posting actually mentions
 - skills marked "used" go on the Languages / Tools lines, "course" skills go on a Coursework line
 - reorders the two project blocks by relevance to the posting"""
import re, shutil
from pathlib import Path
from copy import deepcopy
from docx import Document
from .skills_catalog import CATALOG

ALIAS = {"JavaScript": ["js", "javascript", "node"], "TypeScript": ["ts", "typescript"], "Go": ["golang"], "Kubernetes": ["k8s"],
         "PostgreSQL": ["postgres", "postgresql"], "C++": ["cpp", "c\\+\\+"], "C#": ["c#", "csharp"], "CI/CD": ["ci/cd", "cicd", "continuous integration"],
         "REST APIs": ["rest", "restful"], "LLMs": ["llm", "llms", "large language model"], "Machine Learning": ["machine learning", "ml"],
         "Deep Learning": ["deep learning"], "NLP": ["nlp", "natural language"], "RAG": ["rag", "retrieval[- ]augmented"], "Computer Vision": ["computer vision"],
         "Data Structures": ["data structures"], "Algorithms": ["algorithms"], "Operating Systems": ["operating systems"], "Distributed Systems": ["distributed systems"],
         "AWS EC2": ["ec2"], "AWS S3": ["s3"], "AWS Lambda": ["lambda"], "Terraform": ["terraform"], "OOP": ["object[- ]oriented", "oop"],
         "Unit Testing": ["unit test(ing|s)?"], "System Design": ["system design"], "Microservices": ["microservices?"], "GraphQL": ["graphql"]}
AMBIGUOUS = {"Go", "R", "C", "Make", "Cursor", "Ray", "Vault", "Gin", "Hashing", "Debugging", "Monitoring", "Notion", "Optimization", "Statistics"}
LANG = set(CATALOG["Languages"].split("|"))
CAT_OF = {s: c for c, v in CATALOG.items() for s in v.split("|")}


def _pat(skill):
    alts = [re.escape(skill.lower())] + [a if "\\" in a or "[" in a else re.escape(a) for a in ALIAS.get(skill, [])]
    return re.compile(r"(?<![\w+#])(" + "|".join(alts) + r")(?![\w#]|\+\+)", re.I)


def matched_skills(text: str, bank: dict) -> list[tuple[str, str]]:
    out = []
    for skill, level in bank.items():
        if skill in AMBIGUOUS and skill not in ALIAS:
            continue
        if skill == "Go":   # only when written as a language: "golang", or Go sitting in a list of other technologies
            T = r"(?:Python|Java|Rust|C\+\+|Kubernetes|Docker|TypeScript|Kotlin|Scala|Node)"
            if re.search(r"\bgolang\b", text, re.I) or re.search(rf"\b{T}\b[,/ ]+(?:and |or )?Go\b|\bGo\b[,/ ]+(?:and |or )?{T}\b", text):
                out.append((skill, level))
            continue
        if _pat(skill).search(text):
            out.append((skill, level))
    return out


def _has(line: str, skill: str) -> bool:
    return re.search(r"(?<![\w+#])" + re.escape(skill) + r"(?![\w#]|\+\+)", line, re.I) is not None


def _para_idx(doc, pred, start=0):
    for i in range(start, len(doc.paragraphs)):
        if pred(doc.paragraphs[i].text):
            return i
    return None


def _append_to_list(par, items, before_bold=None):
    """Append ', a, b' to the list in a skills paragraph, ahead of a following bold label such as 'Interests'."""
    if not items:
        return
    runs = par.runs
    target = runs[-1]
    if before_bold:
        for i, r in enumerate(runs):
            if r.bold and r.text.strip().startswith(before_bold):
                j = i - 1
                while j > 0 and not runs[j].text.strip():
                    j -= 1
                target = runs[j]
                break
    target.text = target.text.rstrip() + ", " + ", ".join(items)


def _project_blocks(doc):
    ps = doc.paragraphs
    a = _para_idx(doc, lambda t: t.strip().upper() == "PROJECTS")
    b = _para_idx(doc, lambda t: t.strip().upper() == "EXPERIENCE")
    if a is None or b is None:
        return []
    blocks, cur = [], None
    for i in range(a + 1, b):
        p = ps[i]
        if p.runs and p.runs[0].bold and p.text.strip():
            cur = [p]; blocks.append(cur)
        elif cur is not None and p.text.strip():
            cur.append(p)
    return blocks


def tailor(src_docx, job_text, bank, out_path, max_add=8, max_course=12):
    """Returns dict(info). Writes out_path."""
    doc = Document(str(src_docx))
    ms = matched_skills(job_text, bank)
    lang_line = _para_idx(doc, lambda t: t.startswith("Programming Languages"))
    tool_line = _para_idx(doc, lambda t: t.startswith("Frameworks & Tools"))
    added_used, added_course = [], []
    if lang_line is not None and tool_line is not None:
        lt, tt = doc.paragraphs[lang_line].text, doc.paragraphs[tool_line].text
        for skill, level in ms:
            if _has(lt, skill) or _has(tt, skill):
                continue
            (added_used if level == "used" else added_course).append(skill)
        # "used" skills go on the natural line; "course" skills go on a coursework line
        langs = [s for s in added_used if s in LANG][:max_add]
        tools = [s for s in added_used if s not in LANG][:max_add]
        _append_to_list(doc.paragraphs[lang_line], langs)
        _append_to_list(doc.paragraphs[tool_line], tools, before_bold="Interests")
        course = added_course[:max_course]
        if course:
            src = doc.paragraphs[lang_line]
            new = deepcopy(src._p)
            src._p.addnext(new) if False else doc.paragraphs[tool_line]._p.addnext(new)
            from docx.text.paragraph import Paragraph
            np_ = Paragraph(new, src._parent)
            for r in np_.runs[2:]:
                r._r.getparent().remove(r._r)
            np_.runs[0].text = "Coursework & Familiarity"
            np_.runs[1].text = ": " + ", ".join(course)
        added_used, added_course = langs + tools, course
    # project order by relevance (count of matched skill names inside each block)
    blocks = _project_blocks(doc)
    reordered = False
    if len(blocks) == 2:
        score = lambda bl: sum(1 for s, _ in ms if _has(" ".join(p.text for p in bl), s))
        if score(blocks[1]) > score(blocks[0]):
            first = blocks[0][0]._p
            for p in blocks[1]:
                first.addprevious(p._p)
            reordered = True
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"matched": [s for s, _ in ms], "added_used": added_used, "added_course": added_course, "reordered": reordered}
