"""Fail-closed resume artifact and browser attachment verification."""
from __future__ import annotations

import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_RESUME_BYTES = 10 * 1024 * 1024
MIN_EXTRACTED_CHARS = 120


@dataclass(frozen=True)
class ResumeArtifact:
    path: str
    name: str
    size: int
    sha256: str
    file_type: str
    pages: int | None
    extracted_chars: int
    warnings: tuple[str, ...] = ()

    def public(self) -> dict:
        return asdict(self)


class ResumeValidationError(ValueError):
    pass


def _normalized(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", str(value).lower()))


def _extract_pdf(path: Path) -> tuple[str, int]:
    if not path.read_bytes()[:5] == b"%PDF-":
        raise ResumeValidationError("file does not have a PDF header")
    try:
        reader = PdfReader(str(path), strict=True)
    except Exception as exc:
        raise ResumeValidationError(f"PDF cannot be opened: {str(exc)[:160]}") from exc
    if reader.is_encrypted:
        raise ResumeValidationError("PDF is encrypted")
    pages = len(reader.pages)
    if not 1 <= pages <= 2:
        raise ResumeValidationError(f"resume must be 1 or 2 pages, found {pages}")
    text = []
    for index, page in enumerate(reader.pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if not (400 <= width <= 700 and 550 <= height <= 1000):
            raise ResumeValidationError(f"page {index} has unusual dimensions ({width:.0f} x {height:.0f} points)")
        try:
            text.append(page.extract_text() or "")
        except Exception as exc:
            raise ResumeValidationError(f"text extraction failed on page {index}: {str(exc)[:120]}") from exc
    return "\n".join(text), pages


def _extract_docx(path: Path) -> str:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ResumeValidationError(f"DOCX cannot be opened: {str(exc)[:160]}") from exc
    parts = [p.text for p in doc.paragraphs]
    parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def validate_resume_artifact(path: str | Path, profile: dict) -> ResumeArtifact:
    """Require a readable resume containing the verified applicant identity."""
    if not path:
        raise ResumeValidationError("no resume file was selected")
    resume = Path(path).resolve()
    if not resume.is_file():
        raise ResumeValidationError("resume file does not exist")
    size = resume.stat().st_size
    if not 100 <= size <= MAX_RESUME_BYTES:
        raise ResumeValidationError(f"resume size is invalid ({size} bytes)")
    suffix = resume.suffix.lower()
    if suffix == ".pdf":
        text, pages = _extract_pdf(resume)
    elif suffix == ".docx":
        text, pages = _extract_docx(resume), None
    elif suffix == ".txt":
        text, pages = resume.read_text(errors="replace"), None
    else:
        raise ResumeValidationError(f"unsupported submission format: {suffix or 'none'}")
    compact = re.sub(r"\s+", "", text)
    if len(compact) < MIN_EXTRACTED_CHARS:
        raise ResumeValidationError(f"resume has too little extractable text ({len(compact)} characters)")

    identity = profile.get("identity") or {}
    education = profile.get("education") or {}
    normalized_text = _normalized(text)
    missing = []
    for label, expected in (
        ("first name", identity.get("first_name")),
        ("last name", identity.get("last_name")),
        ("school", education.get("school")),
    ):
        if expected and _normalized(expected) not in normalized_text:
            missing.append(label)
    if missing:
        raise ResumeValidationError("resume does not contain the verified " + ", ".join(missing))
    warnings = []
    email = str(identity.get("email") or "").lower()
    if email and re.sub(r"\s+", "", text.lower()).find(email) < 0:
        warnings.append("resume email differs from the application profile")

    return ResumeArtifact(
        path=str(resume), name=resume.name, size=size,
        sha256=hashlib.sha256(resume.read_bytes()).hexdigest(),
        file_type=suffix.lstrip("."), pages=pages,
        extracted_chars=len(compact), warnings=tuple(warnings),
    )


FILE_METADATA_JS = """async el => {
  const file = el && el.files && el.files[0];
  if (!file) return {error: 'file input is empty'};
  const bytes = new Uint8Array(await file.arrayBuffer());
  let sha256;
  if (globalThis.crypto && globalThis.crypto.subtle) {
    const digest = await globalThis.crypto.subtle.digest('SHA-256', bytes);
    sha256 = [...new Uint8Array(digest)].map(x => x.toString(16).padStart(2, '0')).join('');
  } else {
    const rotr = (x, n) => (x >>> n) | (x << (32 - n));
    const k = [
      0x428a2f98,0x71374491,0xb5c0fbcf,0xe9b5dba5,0x3956c25b,0x59f111f1,0x923f82a4,0xab1c5ed5,
      0xd807aa98,0x12835b01,0x243185be,0x550c7dc3,0x72be5d74,0x80deb1fe,0x9bdc06a7,0xc19bf174,
      0xe49b69c1,0xefbe4786,0x0fc19dc6,0x240ca1cc,0x2de92c6f,0x4a7484aa,0x5cb0a9dc,0x76f988da,
      0x983e5152,0xa831c66d,0xb00327c8,0xbf597fc7,0xc6e00bf3,0xd5a79147,0x06ca6351,0x14292967,
      0x27b70a85,0x2e1b2138,0x4d2c6dfc,0x53380d13,0x650a7354,0x766a0abb,0x81c2c92e,0x92722c85,
      0xa2bfe8a1,0xa81a664b,0xc24b8b70,0xc76c51a3,0xd192e819,0xd6990624,0xf40e3585,0x106aa070,
      0x19a4c116,0x1e376c08,0x2748774c,0x34b0bcb5,0x391c0cb3,0x4ed8aa4a,0x5b9cca4f,0x682e6ff3,
      0x748f82ee,0x78a5636f,0x84c87814,0x8cc70208,0x90befffa,0xa4506ceb,0xbef9a3f7,0xc67178f2];
    const bitLength = bytes.length * 8;
    const paddedLength = Math.ceil((bytes.length + 9) / 64) * 64;
    const padded = new Uint8Array(paddedLength); padded.set(bytes); padded[bytes.length] = 0x80;
    const view = new DataView(padded.buffer); view.setUint32(paddedLength - 4, bitLength >>> 0, false);
    let h = [0x6a09e667,0xbb67ae85,0x3c6ef372,0xa54ff53a,0x510e527f,0x9b05688c,0x1f83d9ab,0x5be0cd19];
    const w = new Uint32Array(64);
    for (let offset = 0; offset < paddedLength; offset += 64) {
      for (let i = 0; i < 16; i++) w[i] = view.getUint32(offset + i * 4, false);
      for (let i = 16; i < 64; i++) {
        const a = rotr(w[i-15],7) ^ rotr(w[i-15],18) ^ (w[i-15] >>> 3);
        const b = rotr(w[i-2],17) ^ rotr(w[i-2],19) ^ (w[i-2] >>> 10);
        w[i] = (w[i-16] + a + w[i-7] + b) >>> 0;
      }
      let [a,b,c,d,e,f,g,hh] = h;
      for (let i = 0; i < 64; i++) {
        const s1 = rotr(e,6) ^ rotr(e,11) ^ rotr(e,25), ch = (e & f) ^ (~e & g);
        const t1 = (hh + s1 + ch + k[i] + w[i]) >>> 0;
        const s0 = rotr(a,2) ^ rotr(a,13) ^ rotr(a,22), maj = (a & b) ^ (a & c) ^ (b & c);
        const t2 = (s0 + maj) >>> 0;
        hh=g; g=f; f=e; e=(d+t1)>>>0; d=c; c=b; b=a; a=(t1+t2)>>>0;
      }
      h = h.map((x,i) => (x + [a,b,c,d,e,f,g,hh][i]) >>> 0);
    }
    sha256 = h.map(x => x.toString(16).padStart(8, '0')).join('');
  }
  return {name: file.name, size: file.size, type: file.type, sha256};
}"""


def browser_file_metadata(page, field_id: str) -> dict | None:
    """Return the exact File object metadata, or None if a previous-step input is gone."""
    locator = page.locator(f'[data-ja="{field_id}"]')
    if not locator.count():
        return None
    return locator.first.evaluate(FILE_METADATA_JS)


def verify_browser_file(metadata: dict | None, artifact: ResumeArtifact) -> dict:
    """Compare the employer form's File object with the validated local artifact."""
    if not metadata:
        raise ResumeValidationError("resume file input could not be read after upload")
    if metadata.get("error"):
        raise ResumeValidationError(str(metadata["error"]))
    mismatches = []
    if metadata.get("name") != artifact.name:
        mismatches.append(f"filename {metadata.get('name')!r} != {artifact.name!r}")
    if metadata.get("size") != artifact.size:
        mismatches.append(f"size {metadata.get('size')} != {artifact.size}")
    if metadata.get("sha256") != artifact.sha256:
        mismatches.append("SHA-256 fingerprint differs")
    if mismatches:
        raise ResumeValidationError("; ".join(mismatches))
    return {
        "status": "verified",
        "name": artifact.name,
        "size": artifact.size,
        "sha256": artifact.sha256,
        "method": "browser-file-byte-match",
    }
